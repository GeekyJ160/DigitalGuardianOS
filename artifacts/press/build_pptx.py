#!/usr/bin/env python3
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

ROOT = Path("/workspace/artifacts/press")
OUT = Path("/workspace/artifacts/GuardianOS-campaign.pptx")
PUBLIC = Path("/workspace/public/campaign")

NOTES = [
    "Open on the black box. GuardianOS is a personal safety and evidence OS — not another SOS button. Tagline: Stay protected. Preserve the truth.",
    "Name the moment: late meetings, first dates, rideshares. If something happens, people need the record, not guesses.",
    "Position against trackers and panic apps. The category is Prevent, Detect, Preserve, Escalate, Reconstruct.",
    "Walk the five verbs. Every protected session writes an original, encrypted Guardian Capsule.",
    "The product sentence: authorize a session with a check-in time. Guardian records facts and follows the user’s protocol. It does not decide danger.",
    "The Capsule holds breadcrumbs, check-ins, notes, and integrity hashes. Language is observed event — never a verdict.",
    "Circle and protocol. Missed check-ins escalate only as the user authorized, including dead-man evidence escrow.",
    "Covert layer: PIN, phrase, gesture; decoy calculator; escrow if the user cannot act.",
    "Close on the tagline. Call to action: start a Guardian Session. Write your protocol. Keep the record original.",
]

prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

for i in range(1, 10):
    img = ROOT / "slides" / f"s{i}.png"
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(
        str(img), Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height
    )
    slide.notes_slide.notes_text_frame.text = NOTES[i - 1]

prs.core_properties.title = "GuardianOS — Campaign briefing"
prs.core_properties.author = "GuardianOS"
prs.core_properties.subject = "Stay protected. Preserve the truth."
prs.core_properties.keywords = "GuardianOS, safety, evidence, capsule, campaign"

PUBLIC.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
prs.save(PUBLIC / "GuardianOS-campaign.pptx")
print("wrote", OUT)
